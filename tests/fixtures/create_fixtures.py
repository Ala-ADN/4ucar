"""Script to generate test fixture files. Run once before tests if fixtures are missing."""

from pathlib import Path

HERE = Path(__file__).parent


def create_clean_excel():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Données académiques"
    ws.append(["Taux de réussite", "Taux d'abandon", "Taux de redoublement", "Effectif total"])
    ws.append(["75.5", "10.2", "14.3", "1200"])
    ws.append(["82", "8", "10", "980"])
    ws.append(["68", "15", "17", "750"])
    wb.save(HERE / "clean_academic.xlsx")


def create_merged_cells_excel():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feuille1"
    # Merge header cells to simulate real-world admin files
    ws.merge_cells("A1:B1")
    ws["A1"] = "Données financières"
    ws.append(["Budget alloué", "Budget consommé", "Taux d'exécution"])
    ws.append(["500000", "420000", "84"])
    ws.append(["300000", "280000", "93.3"])
    wb.save(HERE / "merged_cells_finance.xlsx")


def create_semicolon_csv():
    content = "Taux de réussite;Taux d'abandon;Taux de redoublement;Effectif total\n"
    content += "75,5;10,2;14,3;1200\n"
    content += "82;8;10;980\n"
    (HERE / "semicolon_academic.csv").write_text(content, encoding="utf-8")


def create_comma_csv():
    content = "success_rate,dropout_rate,enrollment_total\n"
    content += "75.5,10.2,1200\n"
    content += "82,8,980\n"
    (HERE / "comma_academic.csv").write_text(content, encoding="utf-8")


def create_windows1252_csv():
    content = "Taux de réussite;Effectif total\n75,5;1200\n"
    (HERE / "windows1252_academic.csv").write_bytes(content.encode("windows-1252"))


if __name__ == "__main__":
    create_clean_excel()
    create_merged_cells_excel()
    create_semicolon_csv()
    create_comma_csv()
    create_windows1252_csv()
    print("Fixtures created.")
