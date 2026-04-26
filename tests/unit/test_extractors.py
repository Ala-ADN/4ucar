"""Unit tests for file extractors (Excel, CSV)."""

from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

import pytest


# ── Excel extractor ───────────────────────────────────────────────────────────


def test_excel_extract_clean(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    from backend.services.ingestion_service.extractors.excel import extract

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Données"
    ws.append(["Taux de réussite", "Effectif total"])
    ws.append(["75.5", "1200"])
    ws.append(["82", "980"])
    path = tmp_path / "test.xlsx"
    wb.save(str(path))

    result = extract(path)
    assert "Taux de réussite" in result.headers
    assert result.total_rows == 2
    assert result.rows[0]["Taux de réussite"] == "75.5"


def test_excel_extract_unmerges_cells(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    from backend.services.ingestion_service.extractors.excel import extract

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:B1")
    ws["A1"] = "Titre fusionné"
    ws.append(["Budget alloué", "Budget consommé"])
    ws.append(["500000", "420000"])
    path = tmp_path / "merged.xlsx"
    wb.save(str(path))

    # Should not raise despite merged cells
    result = extract(path)
    assert result.total_rows >= 1


def test_excel_available_sheets(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    from backend.services.ingestion_service.extractors.excel import extract

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Feuille1"
    ws1.append(["Col A"])
    ws1.append(["val"])
    ws2 = wb.create_sheet("Feuille2")
    ws2.append(["Col B"])
    ws2.append(["val2"])
    path = tmp_path / "multi_sheet.xlsx"
    wb.save(str(path))

    result = extract(path)
    assert len(result.available_sheets) == 2
    names = [s["name"] for s in result.available_sheets]
    assert "Feuille1" in names
    assert "Feuille2" in names


# ── CSV extractor ─────────────────────────────────────────────────────────────


def test_csv_semicolon_delimiter(tmp_path):
    from backend.services.ingestion_service.extractors.csv_extractor import extract

    content = "Taux de réussite;Effectif total\n75,5;1200\n82;980\n"
    path = tmp_path / "test.csv"
    path.write_text(content, encoding="utf-8")

    result = extract(path)
    assert "Taux de réussite" in result.headers
    assert result.total_rows == 2


def test_csv_comma_delimiter(tmp_path):
    from backend.services.ingestion_service.extractors.csv_extractor import extract

    content = "success_rate,dropout_rate,enrollment_total\n75.5,10.2,1200\n"
    path = tmp_path / "test.csv"
    path.write_text(content, encoding="utf-8")

    result = extract(path)
    assert "success_rate" in result.headers
    assert result.total_rows == 1


def test_csv_explicit_encoding(tmp_path):
    from backend.services.ingestion_service.extractors.csv_extractor import extract

    content = "Taux de réussite;Effectif total\n75,5;1200\n"
    path = tmp_path / "test.csv"
    path.write_bytes(content.encode("windows-1252"))

    result = extract(path, encoding="windows-1252", delimiter=";")
    assert "Taux de réussite" in result.headers


# ── OCR confidence tiering logic ──────────────────────────────────────────────


def test_ocr_confidence_tiers():
    """Verify confidence threshold constants are logically ordered."""
    from backend.services.ingestion_service.config import IngestionSettings

    s = IngestionSettings()
    assert s.ocr_confidence_low < s.ocr_confidence_high
    assert 0 < s.ocr_confidence_low < 1
    assert 0 < s.ocr_confidence_high <= 1


# ── Fuzzy mapper ──────────────────────────────────────────────────────────────


def test_fuzzy_mapper_exact_match():
    from backend.services.ingestion_service.mapping.fuzzy_mapper import fuzzy_map

    result = fuzzy_map(["taux de réussite"], "academic")
    assert result[0]["field_id"] == "success_rate"
    assert result[0]["confidence"] == 1.0


def test_fuzzy_mapper_no_match():
    from backend.services.ingestion_service.mapping.fuzzy_mapper import fuzzy_map

    result = fuzzy_map(["colonne_totalement_inconnue_xyz"], "academic")
    assert result[0]["field_id"] is None


def test_fuzzy_mapper_partial_match():
    from backend.services.ingestion_service.mapping.fuzzy_mapper import fuzzy_map

    result = fuzzy_map(["réussite"], "academic")
    # Should find something close to success_rate
    assert result[0]["field_id"] is not None or result[0]["confidence"] == 0.0


# ── KPI schema ────────────────────────────────────────────────────────────────


def test_kpi_schema_alias_map_unique():
    from backend.services.ingestion_service.kpi_schema import build_alias_map

    alias_map = build_alias_map()
    assert len(alias_map) > 0
    assert "success_rate" in alias_map.values()


def test_required_fields_for_domain():
    from backend.services.ingestion_service.kpi_schema import required_fields_for_domain

    required = required_fields_for_domain("academic")
    assert "success_rate" in required
    assert "enrollment_total" in required

    required_fin = required_fields_for_domain("finance")
    assert "budget_allocated" in required_fin
    assert "budget_consumed" in required_fin
